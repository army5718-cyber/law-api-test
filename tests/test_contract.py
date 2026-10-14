import re, glob, zipfile, sys, collections, io
import openpyxl
from xml.dom import minidom
XLSM=sys.argv[1]; SRC=sys.argv[2]
FAILS=[]; OK=[0]
def check(c,m):
    if c: OK[0]+=1
    else: FAILS.append(m); print('FAIL:',m)
def strip(line):
    out=[];ins=False;i=0
    while i<len(line):
        ch=line[i]
        if ch=='"':
            if ins and i+1<len(line) and line[i+1]=='"': out.append('  '); i+=2; continue
            ins=not ins; out.append('"'); i+=1; continue
        if ch=="'" and not ins: break
        out.append(ch); i+=1
    return ''.join(out)
def logical(txt):
    res=[];buf=''
    for l in txt.splitlines():
        s=strip(l)
        if s.rstrip().endswith(' _'): buf+=s.rstrip()[:-1]+' '; continue
        res.append(buf+s); buf=''
    return res
files={f.split('/')[-1]:open(f,encoding='utf8').read() for f in [x for e in ('bas','cls','frm') for x in glob.glob(SRC+'/*.'+e)]}
# ---------- 프로시저 시그니처
procs=collections.defaultdict(list)   # name -> [(module,public,kind,minargs,maxargs)]
for fn,txt in files.items():
    mod=re.search(r'VB_Name = "([^"]+)"',txt).group(1)
    for ln in logical(txt):
        m=re.match(r'\s*(Public |Private |Friend )?(Static )?(Sub|Function)\s+([\w가-힣]+)\s*(\((.*)\))?',ln)
        if m:
            params=[p for p in (m.group(6) or '').split(',') if p.strip()] if m.group(6) is not None else []
            # 괄호 안 배열 표기 () 때문에 split 이 깨지지 않도록 처리
            raw=m.group(6) or ''
            params=[]; depth=0; cur=''
            for ch in raw:
                if ch=='(': depth+=1
                if ch==')': depth-=1
                if ch==',' and depth==0: params.append(cur); cur=''
                else: cur+=ch
            if cur.strip(): params.append(cur)
            mn=sum(1 for p in params if not re.match(r'\s*Optional\b',p)); mx=len(params)
            procs[m.group(4)].append((mod,(m.group(1) or 'Public ').strip()!='Private',m.group(3),mn,mx,fn))
public={n:v for n,v in procs.items()}
# ---------- 호출 인자 수 점검
def top_args(s):
    args=[];depth=0;cur='';ins=False
    for ch in s:
        if ch=='"': ins=not ins
        if not ins:
            if ch=='(' : depth+=1
            if ch==')' : depth-=1
        if ch==',' and depth==0 and not ins: args.append(cur); cur=''
        else: cur+=ch
    if cur.strip(): args.append(cur)
    return args
calls=0
for fn,txt in files.items():
    mod=re.search(r'VB_Name = "([^"]+)"',txt).group(1)
    for ln in logical(txt):
        t=ln.strip()
        if re.match(r'(Public |Private |Friend )?(Static )?(Sub|Function|Dim|Const|Private|Public|Attribute|Option|End|Declare)',t): 
            if not re.match(r'(Public |Private |Friend )?(Static )?(Sub|Function)\b',t): pass
        # 문자열 제거본
        u=re.sub(r'"[^"]*"','""',t)
        for name,defs in public.items():
            if len(name)<3: continue
            for m in re.finditer(r'(?<![\w가-힣.])'+re.escape(name)+r'(?![\w가-힣])',u):
                if re.match(r'\s*(Public |Private |Friend )?(Static )?(Sub|Function)\s+'+re.escape(name),u): continue
                if re.search(r'\bEnd\s+(Sub|Function)',u): continue
                rest=u[m.end():]
                # 같은 모듈 private 또는 public 만 호출 가능
                cand=[d for d in defs if d[1] or d[0]==mod]
                if not cand: 
                    if re.match(r'\s*[\(\s,]',rest) or rest.strip()=='' : check(False,f'{mod}: {name} 는 다른 모듈의 Private 인데 호출됨 ({t[:60]})')
                    continue
                # 호출 형태: Name(args) 또는 Name args (Call 문)
                if rest.startswith('('):
                    depth=0; end=None
                    for i,ch in enumerate(rest):
                        if ch=='(': depth+=1
                        if ch==')':
                            depth-=1
                            if depth==0: end=i; break
                    inner=rest[1:end] if end else ''
                    n=len(top_args(inner)) if inner.strip() else 0
                    # 프로시저 호출 문(Sub 이름 뒤 괄호)이거나 함수 호출
                    ok=any(d[3]<=n<=d[4] for d in cand)
                    calls+=1
                    check(ok,f'{mod}: {name} 인자 수 {n} (허용 {[ (d[3],d[4]) for d in cand]}) -> {t[:70]}')
                else:
                    stm=re.match(r'^(?:Call\s+)?'+re.escape(name)+r'(?![\w가-힣])\s*(.*)$',u.strip())
                    if stm:   # 문장 맨 앞의 프로시저 호출
                        arg=stm.group(1)
                        if arg.startswith('=') : continue
                        n=len(top_args(arg)) if arg.strip() else 0
                        ok=any(d[3]<=n<=d[4] for d in cand); calls+=1
                        check(ok,f'{mod}: {name} (문장) 인자 수 {n} (허용 {[ (d[3],d[4]) for d in cand]}) -> {t[:70]}')
print('호출 점검',calls)
# ---------- OnAction 매크로 존재
acts=set()
for fn,txt in files.items():
    for m in re.finditer(r'macros? = Array\(([^)]*)\)',txt): acts|=set(re.findall(r'"(\w+)"',m.group(1)))
    acts|=set(re.findall(r'\.OnAction = "(\w+)"',txt))
z=zipfile.ZipFile(XLSM)
dx=z.read('xl/drawings/drawing1.xml').decode('utf8'); minidom.parseString(dx)
acts|=set(re.findall(r'macro="(\w+)"',dx))
for a in sorted(acts):
    d=public.get(a); check(bool(d) and any(x[1] and x[2]=='Sub' and x[3]==0 for x in d),f'OnAction 매크로 {a} 없음/인자있음')
print('매크로',sorted(acts))
# ---------- 이름 정의 사용 점검
wbx=z.read('xl/workbook.xml').decode('utf8')
defined=set(re.findall(r'<definedName name="([^"]+)"',wbx))
used=set()
for txt in files.values():
    used|=set(re.findall(r'names\("([^"]+)"\)',txt)); used|=set(re.findall(r'Setting(?:Text|Long|Time|Raw)\("([^"]+)"\)',txt)); used|=set(re.findall(r'NameExists\("([^"]+)"\)',txt))
    used|=set(re.findall(r'Formula1:="=([^"]+)"',txt))
for n in sorted(used-{'Print_Area'}): check(n in defined,f'이름 정의 {n} 없음')
print('이름',sorted(used))
# ---------- 시트/표/열 점검
wb=openpyxl.load_workbook(XLSM,keep_vba=True)
consts={}
for txt in files.values():
    for m in re.finditer(r'Public Const (\w+) As (?:String|Long) = ("?[^\r\n\']+"?)',txt): consts[m.group(1)]=m.group(2).strip().strip('"')
for k,v in consts.items():
    if k.startswith('SH_'): check(v in wb.sheetnames,f'시트 {v} 없음')
raw=wb[consts['SH_RAW']]
check('tblRaw' in raw.tables,'tblRaw 없음')
t=raw.tables['tblRaw']; heads=[c.name for c in t.tableColumns]
exp=['접수일자','구분','권역','상호','용량','지역','시간','수수료','검사자','비고','관리번호']
check(heads==exp,f'tblRaw 열 {heads}')
rw={k:int(v) for k,v in consts.items() if k.startswith('RW_') and v.isdigit()}
for k,i in zip(['DATE','KIND','ZONE','NAME','CAP','REG','TIME','FEE','INSP','NOTE','ID'],range(1,12)): check(rw['RW_'+k]==i,f'RW_{k}={rw["RW_"+k]}')
check(raw.cell(1,rw['RW_INSP']).value=='검사자' and raw.cell(1,rw['RW_ID']).value=='관리번호','RW 헤더 위치')
ros=wb[consts['SH_ROSTER']]
check([ros.cell(1,c).value for c in (1,2,3)]==['검사자명','상태','오늘근무'],'01 머리글')
check(ros['H1'].value and ros['I1'].value,'01 별칭 머리글 H/I')
mp=wb[consts['SH_MAP']]
check([mp.cell(1,c).value for c in (1,2,3,4)]==['지역(접수대장 표기)','표준 지역명','권역','소속 클러스터(읍/면/동/구)'],'04 열 구성')
res=wb[consts['SH_RESULT']]
rc={k:int(v) for k,v in consts.items() if k.startswith('R_') and v.isdigit()}
check(rc['R_COLS']==16 and rc['R_DROW']==16 and rc['R_DCOL']==15,'09 열 상수')
check(res.cell(1,rc['R_NAME']).value=='상호' and res.cell(1,rc['R_INSP']).value=='검사자' and res.cell(1,rc['R_TYPE']).value=='유형','09 머리글 위치')
for c in (rc['R_CAP'],rc['R_REGION'],rc['R_DUR'],rc['R_TIME'],rc['R_START']): check(res.cell(5,c).number_format=='@',f'09 텍스트 서식 열{c}')
log=wb[consts['SH_LOG']]; check([log.cell(1,c).value for c in range(1,6)]==['일시','유형','내용','관련업무(상호)','처리상태'],'07 머리글')
st=wb[consts['SH_SETTINGS']]
vals={st.cell(r,2).value:st.cell(r,3).value for r in range(4,16)}
for k in ['운영시작시간','운영종료시간','슬롯단위_분','정식설비기준용량_kW','정식설비소요시간_분','약식설비소요시간_분','온라인업무소요시간_분','점심시간길이_분','점심탐색순서','표시시작시간','표시종료시간']:
    check(vals.get(k) not in (None,''),f'02 설정값 {k}')
check(str(vals['표시종료시간'])=='24:00' and str(vals['표시시작시간'])=='00:00','02 표시시간 24h')
day=wb[consts['SH_DAY']]
check(day.row_dimensions[3].hidden,'11 3행 숨김')
cd=day.column_dimensions; check(cd['M'].hidden and cd['M'].min<=13 and cd['M'].max>=14,'11 M:N 숨김')
mon=wb[consts['SH_MON']]; check(mon['C3'].value is not None,'12 C3 기준월')
# 14열(M,N) 및 GRID_COL0
check(int(consts['GRID_COL0'])==15 and int(consts['SCR_NOTE_L'])==13 and int(consts['SCR_NOTE_R'])==14,'화면 열 상수')
# 정적 화면 구조가 VBA 기대와 일치: A열 병합 라벨, 머리글 6행, 현황판 머리글 행
check(day['A5'].value=='정기검사' and day['H5'].value=='사용전검사','11 5행 제목')
check(day.cell(6,15).value=='시간' and day.cell(6,16).value,'11 현황판 머리글(6행)')
labs=[c.value for c in day['A'] if c.value in('부천지역','김포지역','비고')]; check(labs[:2]==['부천지역','김포지역'] and '비고' in labs,f'11 A열 라벨 {labs}')
# ---------- 시트 코드명 <-> 문서 모듈
import olefile
ole=olefile.OleFileIO(io.BytesIO(z.read('xl/vbaProject.bin')))
proj=ole.openstream('PROJECT').read().decode('cp949')
docs=set(re.findall(r'Document=([^/\r\n]+)/',proj))
codes={ws.sheet_properties.codeName for ws in wb}
check(codes<=docs,f'문서 모듈 없는 시트 {codes-docs}')
check(docs-codes-{wb.code_name}==set(),f'시트 없는 문서 모듈 {docs-codes-{wb.code_name}}')
# ---------- VBA 스트림 왕복
sys.path.insert(0,__import__('os').path.dirname(__import__('os').path.abspath(__file__)))
import ovba
mods=[m.strip() for m in re.findall(r'^(?:Module|BaseClass)=(.+)$',proj,re.M)]
for name in list(docs)+[m for m in mods]:
    try: dec=ovba.decompress(ole.openstream('VBA/'+name).read())
    except Exception as e: check(False,f'VBA 스트림 {name}: {e}'); continue
    txt=dec.decode('cp949')
    check(f'Attribute VB_Name = "{name}"' in txt,f'VB_Name {name}')
    if name.startswith('Module'): check('Option Explicit' in txt,f'Option Explicit {name}')
# ---------- 이벤트 시트 모듈이 부르는 프로시저가 Public 인지
for fn,txt in files.items():
    if fn.startswith('Sheet') or fn.startswith('현재'):
        for ln in logical(txt):
            m=re.match(r'\s+([A-Z][\w가-힣]+)(?:\s|$)',ln)
            if m and m.group(1) in public and m.group(1) not in ('Private',): 
                check(any(d[1] for d in public[m.group(1)]),f'{fn}: 비공개 프로시저 호출 {m.group(1)}')
print('OK',OK[0],'FAIL',len(FAILS))
for f in FAILS: print(' -',f)
